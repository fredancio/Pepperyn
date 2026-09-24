"""Local full composition/ASGI proof only; no deployed session or Beta proof."""
import copy
from io import BytesIO
from pathlib import Path

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation

from sandbox.heterogeneous_workbooks import run_recorded_registered_mock_analysis
from services.governed_analysis_persistence import save_governed_analysis, _digest
from services.governed_analysis_read import GovernedReadRefused
from services.governed_output import read_owned_output, export_owned_output
from routers.governed_output import build_governed_output_router
from test_governed_analysis_persistence import _db, ANALYSIS, COMPANY_A, COMPANY_B, ENTITY_A, ENGAGEMENT_A


@pytest.fixture
def stored():
    raw = Path('tests/golden/fixtures/pepperyn_v1_heterogeneous_english.xlsx').read_bytes()
    execution = run_recorded_registered_mock_analysis(raw, 'pepperyn_v1_heterogeneous_english.xlsx')
    db = _db()
    scope = dict(analysis_id=ANALYSIS, company_id=COMPANY_A, entity_id=ENTITY_A, engagement_id=ENGAGEMENT_A)
    save_governed_analysis(db, analysis_row=dict(id=ANALYSIS, company_id=COMPANY_A, entity_id=ENTITY_A,
        source_data_hash=execution.provenance.raw_source_sha256.lower()),
        engagement_id=ENGAGEMENT_A, envelope=execution.analysis.envelope)
    payload = execution.provenance.model_dump(mode='json') | scope
    db.tables['governed_execution_receipts'] = [scope | dict(payload=payload, sha256=_digest(payload))]
    return db, execution


def extract(format, content):
    if format == 'xlsx':
        wb = load_workbook(BytesIO(content))
        return '\n'.join(str(value) for ws in wb for row in ws.iter_rows(values_only=True) for value in row if value is not None)
    if format == 'pdf':
        return '\n'.join(page.extract_text() for page in PdfReader(BytesIO(content)).pages)
    return '\n'.join(shape.text for slide in Presentation(BytesIO(content)).slides
                     for shape in slide.shapes if shape.has_text_frame)


@pytest.mark.parametrize('format', ['xlsx', 'pdf', 'pptx'])
def test_owned_receipt_reaches_all_terminal_exports(stored, format):
    db, execution = stored
    before = copy.deepcopy(db.tables)
    text = extract(format, export_owned_output(db, analysis_id=ANALYSIS, company_id=COMPANY_A, format=format))
    assert ANALYSIS in text and str(execution.provenance.execution_id) in text
    assert execution.provenance.raw_source_sha256 in text
    assert execution.provenance.envelope_sha256 in text
    assert 'Fournisseur simule local' in text and 'Recu durable verifie' in text
    assert 'UNKNOWN' in text and 'Comparabilite financiere non etablie' in text
    assert db.tables == before


@pytest.mark.parametrize('format', ['xlsx', 'pdf', 'pptx'])
def test_legacy_absence_never_becomes_synthetic_attestation(stored, format):
    db, _ = stored
    db.tables['governed_execution_receipts'] = []
    text = extract(format, export_owned_output(db, analysis_id=ANALYSIS, company_id=COMPANY_A, format=format))
    assert 'Origine non attestee' in text
    assert 'Aucun recu durable' in text
    assert 'Fournisseur simule local' not in text
    assert 'Donnees synthetiques' not in text
    assert read_owned_output(db, analysis_id=ANALYSIS, company_id=COMPANY_A)['execution_provenance'] == dict(status='UNATTESTED', receipt=None)


@pytest.mark.parametrize('table', ['governed_execution_receipts', 'decision_feedback'])
def test_outage_refuses_instead_of_missing_claim(stored, monkeypatch, table):
    db, _ = stored
    original = db.from_
    def query(name):
        if name == table: raise RuntimeError('sensitive')
        return original(name)
    monkeypatch.setattr(db, 'from_', query)
    with pytest.raises(GovernedReadRefused, match='^UNAVAILABLE$'):
        export_owned_output(db, analysis_id=ANALYSIS, company_id=COMPANY_A, format='pdf')


@pytest.mark.parametrize('field', ['analysis_id', 'entity_id', 'envelope_sha256', 'raw_source_sha256', 'provider_mode'])
def test_forged_receipt_even_rehashed_refused(stored, field):
    db, _ = stored
    receipt = db.tables['governed_execution_receipts'][0]
    receipt['payload'][field] = 'forged'
    receipt['sha256'] = _digest(receipt['payload'])
    with pytest.raises(GovernedReadRefused, match='UNAVAILABLE'):
        read_owned_output(db, analysis_id=ANALYSIS, company_id=COMPANY_A)


def test_foreign_scope_never_reads_receipt(stored):
    db, _ = stored
    db.log.clear()
    with pytest.raises(GovernedReadRefused, match='NOT_FOUND'):
        read_owned_output(db, analysis_id=ANALYSIS, company_id=COMPANY_B)
    assert not any(x[0] == 'governed_execution_receipts' for x in db.log)


def app_for(db, *, company=COMPANY_A, admitted=True):
    def admission():
        if not admitted: raise HTTPException(403, 'Admission closed')
    app = FastAPI()
    app.include_router(build_governed_output_router(authenticated_company=lambda: company,
        require_admission=admission, database=lambda: db))
    return TestClient(app)


def test_unmounted_transport_contract_with_mock_auth(stored):
    db, execution = stored
    before = copy.deepcopy(db.tables)
    client = app_for(db)
    url = '/api/governed/analyses/' + ANALYSIS
    result = client.get(url + '?company_id=' + COMPANY_B)
    assert result.status_code == 200
    assert result.json()['execution_provenance']['receipt']['execution_id'] == str(execution.provenance.execution_id)
    for format in ('xlsx','pdf','pptx'):
        response = client.get(url + '/export.' + format)
        assert response.status_code == 200
        assert str(execution.provenance.execution_id) in extract(format, response.content)
    assert app_for(db,company=COMPANY_B).get(url+'?company_id='+COMPANY_A).status_code == 404
    assert app_for(db,admitted=False).get(url).status_code == 403
    assert client.post(url).status_code == 405
    assert db.tables == before


def test_new_shared_modules_do_not_import_sandbox():
    import ast
    for name in ('governed_output','governed_export_rendering','temporal_export'):
        tree = ast.parse(Path('services',name+'.py').read_text(encoding='utf-8'))
        assert not any(isinstance(node,ast.ImportFrom) and (node.module or '').startswith('sandbox') for node in ast.walk(tree))
