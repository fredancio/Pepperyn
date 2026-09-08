from __future__ import annotations

import asyncio
import copy

from fastapi import FastAPI, HTTPException
import pytest
from io import BytesIO
from openpyxl import load_workbook
from pypdf import PdfReader

import main
import routers.analyze as analyze
import sandbox.v1_router as v1_routes
from starlette.requests import Request

COMPANY_A = "20000000-0000-0000-0000-000000000001"
COMPANY_B = "20000000-0000-0000-0000-000000000002"
ENTITY_A = "30000000-0000-0000-0000-000000000001"
ENGAGEMENT_A = "40000000-0000-0000-0000-000000000001"


class _Response:
    def __init__(self, data): self.data = data


class _Query:
    def __init__(self, db, table):
        self.db, self.table = db, table
        self.rows = list(db.tables.get(table, ()))
        self.filters = []
    def select(self, _fields): return self
    def eq(self, field, value):
        self.filters.append((field, value))
        self.rows = [row for row in self.rows if row.get(field) == value]
        return self
    def is_(self, field, value):
        assert value == "null"
        return self.eq(field, None)
    def limit(self, count): self.rows = self.rows[:count]; return self
    def upsert(self, row, on_conflict=None):
        assert on_conflict == "report_id,recommendation_id"
        rows = self.db.tables.setdefault(self.table, [])
        existing = next((item for item in rows if item["report_id"] == row["report_id"]
                         and item["recommendation_id"] == row["recommendation_id"]), None)
        if existing is None:
            inserted = copy.deepcopy(row)
            inserted.setdefault("id", f"feedback-{len(rows) + 1}")
            rows.append(inserted)
        else:
            existing.update(copy.deepcopy(row))
        self.rows = [row]
        return self
    def update(self, values):
        self._update_values = copy.deepcopy(values)
        self._is_update = True
        return self
    def execute(self):
        if getattr(self, "_is_update", False):
            matched = []
            for row in self.db.tables.get(self.table, []):
                if all(row.get(field) == value for field, value in self.filters):
                    row.update(copy.deepcopy(self._update_values))
                    matched.append(row)
            self.rows = matched
        return _Response(copy.deepcopy(self.rows))


class _Rpc:
    def __init__(self, db, params): self.db, self.params = db, copy.deepcopy(params)
    def execute(self):
        analysis_row = self.params["p_analysis"]
        envelope_row = self.params["p_envelope"]
        assert analysis_row["mode"] in {"quick", "complete"}
        assert analysis_row["type_document"] in {
            "AUTRE", "BILAN", "BUDGET", "COMMERCIAL", "COMPTE_RESULTAT",
            "INCONNU", "PREVISIONNEL", "TRESORERIE",
        }
        assert any(row["id"] == envelope_row["engagement_id"] and row["entity_id"] == analysis_row["entity_id"]
                   for row in self.db.tables["engagements"])
        self.db.tables.setdefault("analyses", []).append(analysis_row)
        self.db.tables.setdefault("governed_analysis_envelopes", []).append(envelope_row)
        return _Response(analysis_row["id"])


class _Db:
    def __init__(self):
        self.tables = {
            "entities": [{"id": ENTITY_A, "company_id": COMPANY_A, "is_primary": True,
                          "name": "Synthetic", "relation_type": None}],
            "engagements": [{"id": ENGAGEMENT_A, "entity_id": ENTITY_A}],
        }
    def from_(self, table): return _Query(self, table)
    def rpc(self, name, params): assert name == "persist_governed_analysis_v1"; return _Rpc(self, params)


class _FeedbackReadUnavailableDb(_Db):
    def from_(self, table):
        if table == "decision_feedback":
            raise RuntimeError("feedback registry unavailable")
        return super().from_(table)


def _enable(monkeypatch, db, company=COMPANY_A):
    monkeypatch.setenv("ENVIRONMENT", "development")
    monkeypatch.setenv("PEPPERYN_ENABLE_SYNTHETIC_V1_DEMO", "1")
    monkeypatch.setenv("PEPPERYN_SYNTHETIC_V1_COMPANY_ID", COMPANY_A)
    async def auth(authorization, _x_auth_type):
        if not authorization or not authorization.startswith("Bearer "):
            raise HTTPException(status_code=401, detail="Token requis")
        return company, "pro", "admin"
    monkeypatch.setattr(analyze, "_resolve_auth", auth)
    monkeypatch.setattr(main, "get_supabase_service", lambda: db)


def _empty_request() -> Request:
    async def receive(): return {"type": "http.request", "body": b"", "more_body": False}
    return Request({"type": "http", "method": "POST", "path": "/api/v1/synthetic-demo",
                    "headers": []}, receive)


def test_authenticated_no_body_demo_persists_and_reloads_after_cache_free_reconstruction(monkeypatch):
    db = _Db(); _enable(monkeypatch, db)
    created = asyncio.run(v1_routes.run_v1_synthetic_demo(
        request=_empty_request(), authorization="Bearer test", x_auth_type=None,
    ))
    assert created.result.verification_tag == "V1_GOVERNED_SINGLE_CALL"
    assert len(db.tables["governed_analysis_envelopes"]) == 1

    reconstructed = _Db()
    reconstructed.tables = copy.deepcopy(db.tables)
    monkeypatch.setattr(main, "get_supabase_service", lambda: reconstructed)
    loaded = asyncio.run(v1_routes.get_v1_governed_analysis(
        created.analyse_id, authorization="Bearer test", x_auth_type=None,
    ))
    assert loaded.result == created.result
    assert loaded.recommendations_tracking == created.recommendations_tracking


def test_governed_intention_uses_server_snapshot_and_never_confirms_a_decision(monkeypatch):
    import httpx

    db = _Db(); _enable(monkeypatch, db)
    async def exercise():
        app = FastAPI(); app.include_router(v1_routes.router)
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://test",
        ) as client:
            created = await client.post("/api/v1/synthetic-demo", headers={"Authorization": "Bearer test"})
            analysis_id = created.json()["analyse_id"]
            recommendation = created.json()["recommendations_tracking"][0]
            recorded = await client.post(
                f"/api/v1/governed-analyses/{analysis_id}/intention",
                headers={"Authorization": "Bearer test"},
                json={"recommendation_id": recommendation["id"], "status": "planned"},
            )
            forged = await client.post(
                f"/api/v1/governed-analyses/{analysis_id}/intention",
                headers={"Authorization": "Bearer test"},
                json={"recommendation_id": "000000000000", "status": "planned"},
            )
            return recorded, forged, recommendation

    recorded, forged, recommendation = asyncio.run(exercise())
    assert recorded.status_code == 200
    assert recorded.json() == {
        "success": True, "intention_recorded": True,
        "decision_confirmed": False, "arc_created": False,
    }
    assert forged.status_code == 404
    assert len(db.tables["decision_feedback"]) == 1
    row = db.tables["decision_feedback"][0]
    assert row["recommendation_id"] == recommendation["id"]
    assert row["recommendation_text"] == recommendation["text"]
    assert recommendation["rationale"]
    assert recommendation["fact_ids"] or recommendation["prerequisite_validation"]
    assert not db.tables.get("decision_arcs")

    reloaded = asyncio.run(v1_routes.get_v1_governed_analysis(
        row["report_id"], authorization="Bearer test", x_auth_type=None,
    ))
    saved = next(item for item in reloaded.recommendations_tracking
                 if item["id"] == recommendation["id"])
    assert saved["status"] == "planned"
    assert saved["comment"] is None


def test_explicit_decision_requires_intention_and_retains_prerequisites(monkeypatch):
    import httpx

    db = _Db(); _enable(monkeypatch, db)
    async def exercise():
        app = FastAPI(); app.include_router(v1_routes.router)
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
            created = await client.post("/api/v1/synthetic-demo", headers={"Authorization": "Bearer test"})
            analysis_id = created.json()["analyse_id"]
            recommendation = next(
                item for item in created.json()["recommendations_tracking"]
                if item["prerequisite_validation"]
            )
            payload = {
                "recommendation_id": recommendation["id"],
                "decision_kind": "accepted_conditional",
                "decision_text": "Retenir sous réserve des validations listées.",
                "prerequisites_acknowledged": False,
            }
            no_ack = await client.post(
                f"/api/v1/governed-analyses/{analysis_id}/decision",
                headers={"Authorization": "Bearer test"}, json=payload,
            )
            intention = await client.post(
                f"/api/v1/governed-analyses/{analysis_id}/intention",
                headers={"Authorization": "Bearer test"},
                json={"recommendation_id": recommendation["id"], "status": "unsure"},
            )
            payload["prerequisites_acknowledged"] = True
            decision = await client.post(
                f"/api/v1/governed-analyses/{analysis_id}/decision",
                headers={"Authorization": "Bearer test"}, json=payload,
            )
            reload = await client.get(
                f"/api/v1/governed-analyses/{analysis_id}",
                headers={"Authorization": "Bearer test"},
            )
            return no_ack, intention, decision, reload, recommendation

    no_ack, intention, decision, reload, recommendation = asyncio.run(exercise())
    assert no_ack.status_code == 422
    assert intention.status_code == 200
    assert decision.status_code == 200
    assert decision.json() == {
        "success": True, "decision_confirmed": True,
        "decision_confirmation_source": "explicit", "arc_created": False,
    }
    saved = next(item for item in reload.json()["recommendations_tracking"]
                 if item["id"] == recommendation["id"])
    assert saved["status"] == "decided"
    assert saved["decision_kind"] == "accepted_conditional"
    assert saved["decision_confirmation_source"] == "explicit"
    assert saved["prerequisites_acknowledged"] is True
    assert not db.tables.get("decision_arcs")


def test_explicit_decision_without_prior_intention_is_refused(monkeypatch):
    import httpx

    db = _Db(); _enable(monkeypatch, db)
    async def exercise():
        app = FastAPI(); app.include_router(v1_routes.router)
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
            created = await client.post("/api/v1/synthetic-demo", headers={"Authorization": "Bearer test"})
            rec = created.json()["recommendations_tracking"][0]
            return await client.post(
                f"/api/v1/governed-analyses/{created.json()['analyse_id']}/decision",
                headers={"Authorization": "Bearer test"},
                json={"recommendation_id": rec["id"], "decision_kind": "rejected",
                      "decision_text": "Non retenue.", "prerequisites_acknowledged": False},
            )
    response = asyncio.run(exercise())
    assert response.status_code == 409


def test_feedback_read_outage_does_not_hide_verified_governed_analysis(monkeypatch):
    db = _Db(); _enable(monkeypatch, db)
    created = asyncio.run(v1_routes.run_v1_synthetic_demo(
        request=_empty_request(), authorization="Bearer test", x_auth_type=None,
    ))
    unavailable = _FeedbackReadUnavailableDb()
    unavailable.tables = copy.deepcopy(db.tables)
    monkeypatch.setattr(main, "get_supabase_service", lambda: unavailable)

    loaded = asyncio.run(v1_routes.get_v1_governed_analysis(
        created.analyse_id, authorization="Bearer test", x_auth_type=None,
    ))

    assert loaded.result == created.result
    assert loaded.recommendations_tracking
    assert all(item["status"] is None for item in loaded.recommendations_tracking)


def test_feedback_read_outage_refuses_export_instead_of_hiding_a_decision(monkeypatch):
    db = _Db(); _enable(monkeypatch, db)
    created = asyncio.run(v1_routes.run_v1_synthetic_demo(
        request=_empty_request(), authorization="Bearer test", x_auth_type=None,
    ))
    unavailable = _FeedbackReadUnavailableDb()
    unavailable.tables = copy.deepcopy(db.tables)
    monkeypatch.setattr(main, "get_supabase_service", lambda: unavailable)

    with pytest.raises(HTTPException) as error:
        asyncio.run(v1_routes.export_v1_governed_excel(
            created.analyse_id, authorization="Bearer test", x_auth_type=None,
        ))
    assert error.value.status_code == 503
    assert "export gouverné refusé" in error.value.detail


def test_second_company_cannot_reload_first_company_analysis(monkeypatch):
    db = _Db(); _enable(monkeypatch, db)
    created = asyncio.run(v1_routes.run_v1_synthetic_demo(
        request=_empty_request(), authorization="Bearer a", x_auth_type=None,
    ))
    _enable(monkeypatch, db, company=COMPANY_B)
    with pytest.raises(HTTPException) as error:
        asyncio.run(v1_routes.get_v1_governed_analysis(
            created.analyse_id, authorization="Bearer b", x_auth_type=None,
        ))
    assert error.value.status_code == 404


def test_main_application_does_not_mount_demo_without_startup_flag():
    assert not any(route.path.startswith("/api/v1/") for route in main.app.routes)


def test_non_designated_company_cannot_create_synthetic_history(monkeypatch):
    db = _Db(); _enable(monkeypatch, db, company=COMPANY_B)
    with pytest.raises(HTTPException) as error:
        asyncio.run(v1_routes.run_v1_synthetic_demo(
            request=_empty_request(), authorization="Bearer b", x_auth_type=None,
        ))
    assert error.value.status_code == 404
    assert not db.tables.get("analyses")


def test_http_contract_accepts_empty_body_rejects_payload_and_serializes(monkeypatch):
    import httpx

    db = _Db(); _enable(monkeypatch, db)
    async def exercise():
        app = FastAPI()
        app.include_router(v1_routes.router)
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
            missing_auth = await client.post("/api/v1/synthetic-demo")
            with_body = await client.post(
                "/api/v1/synthetic-demo", headers={"Authorization": "Bearer test"}, json={"fixture": "other"},
            )
            created = await client.post(
                "/api/v1/synthetic-demo", headers={"Authorization": "Bearer test"},
            )
            loaded = await client.get(
                f"/api/v1/governed-analyses/{created.json()['analyse_id']}",
                headers={"Authorization": "Bearer test"},
            )
            return missing_auth, with_body, created, loaded
    missing_auth, with_body, created, loaded = asyncio.run(exercise())
    assert missing_auth.status_code == 401
    assert with_body.status_code == 400
    assert created.status_code == loaded.status_code == 200
    assert db.tables["analyses"][0]["mode"] == "complete"
    assert db.tables["analyses"][0]["type_document"] == "AUTRE"
    assert created.json()["result"]["type_document"] == "FINANCIAL_WORKBOOK"
    assert created.json()["result"]["id"] == created.json()["analyse_id"]
    assert loaded.json()["result"] == created.json()["result"]


def test_governed_exports_reload_after_restart_and_preserve_epistemic_labels(monkeypatch):
    import httpx

    db = _Db(); _enable(monkeypatch, db)
    async def exercise():
        app = FastAPI(); app.include_router(v1_routes.router)
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
            created = await client.post("/api/v1/synthetic-demo", headers={"Authorization": "Bearer test"})
            reconstructed = _Db(); reconstructed.tables = copy.deepcopy(db.tables)
            monkeypatch.setattr(main, "get_supabase_service", lambda: reconstructed)
            analysis_id = created.json()["analyse_id"]
            excel = await client.get(f"/api/v1/governed-analyses/{analysis_id}/export.xlsx",
                                     headers={"Authorization": "Bearer test"})
            pdf = await client.get(f"/api/v1/governed-analyses/{analysis_id}/export.pdf",
                                   headers={"Authorization": "Bearer test"})
            pptx = await client.get(f"/api/v1/governed-analyses/{analysis_id}/export.pptx",
                                    headers={"Authorization": "Bearer test"})
            return excel, pdf, pptx
    excel, pdf, pptx = asyncio.run(exercise())
    assert excel.status_code == pdf.status_code == pptx.status_code == 200
    assert excel.content.startswith(b"PK")
    assert pdf.content.startswith(b"%PDF")
    assert pptx.content.startswith(b"PK")

    workbook = load_workbook(BytesIO(excel.content), data_only=False)
    assert workbook.sheetnames == ["Synthese", "Faits sources", "Inferences", "Recommandations", "UNKNOWN"]
    cells = "\n".join(str(cell.value) for sheet in workbook for row in sheet.iter_rows() for cell in row if cell.value is not None)
    assert "Diagnostic (inference)" in cells
    assert "Observation source-matched - severite inferentielle HIGH" in cells
    assert "EBITDA = -145000" in cells
    assert "Validations requises" in cells
    assert "Prerequis" in cells
    assert "Les recommandations IA ne constituent pas des decisions confirmees." in cells
    assert "Faits sources" in workbook.sheetnames

    reader = PdfReader(BytesIO(pdf.content))
    pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    for token in ("Diagnostic - inference", "Faits sources", "Observations gouvernees",
                  "Observation source-matched: EBITDA = -145000", "Severite inferentielle: HIGH",
                  "Inferences et validations",
                  "UNKNOWN et contradictions", "Recommandations proposees",
                  "ne constituent pas des decisions confirmees"):
        assert token in pdf_text


def test_second_company_cannot_export_first_company_analysis(monkeypatch):
    db = _Db(); _enable(monkeypatch, db)
    created = asyncio.run(v1_routes.run_v1_synthetic_demo(
        request=_empty_request(), authorization="Bearer a", x_auth_type=None,
    ))
    _enable(monkeypatch, db, company=COMPANY_B)
    for handler in (v1_routes.export_v1_governed_excel, v1_routes.export_v1_governed_pdf,
                    v1_routes.export_v1_governed_pptx):
        with pytest.raises(HTTPException) as error:
            asyncio.run(handler(created.analyse_id, authorization="Bearer b", x_auth_type=None))
        assert error.value.status_code == 404
