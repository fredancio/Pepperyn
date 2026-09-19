"""Local synthetic HTTP -> owned history -> actual three-format bytes.

In-memory persistence is not live RLS, professional review or browser evidence.
"""
import asyncio
import copy
from io import BytesIO

import httpx
import pytest
from fastapi import FastAPI
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

import sandbox.v1_router as routes
from sandbox.governed_exports import generate_governed_excel, generate_governed_pdf, generate_governed_pptx
from services.governed_analysis_persistence import save_governed_analysis
from services.governed_temporal_continuity import compare_governed_envelopes
from test_v1_multiperiod_persistence import synthetic_envelope
from test_v1_synthetic_routes import _Db, _Query, _enable, COMPANY_A, COMPANY_B


FORMATS = {"xlsx": generate_governed_excel, "pdf": generate_governed_pdf, "pptx": generate_governed_pptx}


def text_of(kind, content):
    if kind == "xlsx":
        wb = load_workbook(BytesIO(content))
        return "\n".join(str(c.value) for ws in wb for row in ws for c in row if c.value is not None)
    if kind == "pdf":
        return "\n".join(p.extract_text() for p in PdfReader(BytesIO(content)).pages)
    return "\n".join(s.text for slide in Presentation(BytesIO(content)).slides for s in slide.shapes if s.has_text_frame)


def save(db, client, period, value, index, company=COMPANY_A):
    eid = f"30000000-0000-0000-0000-{client:012d}"
    eng = f"40000000-0000-0000-0000-{client:012d}"
    aid = f"10000000-0000-0000-0000-{index:012d}"
    if not any(r["id"] == eid for r in db.tables["entities"]):
        db.tables["entities"].append({"id": eid, "company_id": company})
        db.tables["engagements"].append({"id": eng, "entity_id": eid})
    envelope = synthetic_envelope(period, value)
    save_governed_analysis(db, analysis_row={
        "id": aid, "company_id": company, "entity_id": eid,
        "analyse_json": envelope.analysis_result.model_dump(mode="json"),
        "status": "completed", "fichier_nom": "synthetic-parsed-facts",
        "mode": "complete", "type_document": "COMPTE_RESULTAT",
    }, engagement_id=eng, envelope=envelope)
    return aid, envelope


async def get(kind, aid):
    app = FastAPI()
    app.include_router(routes.router)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        return await client.get(f"/api/v1/governed-analyses/{aid}/export.{kind}",
                                headers={"Authorization": "Bearer synthetic"})


@pytest.mark.parametrize("kind", FORMATS)
def test_owned_periods_survive_to_export_with_limits_and_no_writes(monkeypatch, kind):
    db = _Db()
    _enable(monkeypatch, db)
    current, env = save(db, 1, "2026", -40, 12)
    previous, prior = save(db, 1, "2025", -110, 11)
    # Neighbouring clients and foreign tenant must never contribute history.
    other, _ = save(db, 2, "2025", 789123, 21)
    foreign, _ = save(db, 3, "2025", 912345, 31, COMPANY_B)
    before = copy.deepcopy(db.tables)
    response = asyncio.run(get(kind, current))
    assert response.status_code == 200
    text = text_of(kind, response.content)
    for expected in (current, previous, "Comparabilite financiere non etablie",
                     "EBITDA", "-110", "-40", "70", "2025", "2026"):
        assert expected in text
    for e in (env, prior):
        for fact in e.source_facts.facts:
            assert fact.fact_id in text
    assert other not in text and foreign not in text
    assert "789123" not in text and "912345" not in text
    assert db.tables == before
    if kind == "xlsx":
        rows = dict(load_workbook(BytesIO(response.content))["Continuite temporelle"].values)
        assert rows["EBITDA - ecart arithmetique"] == 70  # numeric, not presentation text
    if kind == "pptx":
        slides = Presentation(BytesIO(response.content)).slides
        temporal = [s for s in slides if any("Continuite temporelle" in x.text for x in s.shapes if x.has_text_frame)]
        assert temporal
        assert all(any("Comparabilite financiere non etablie" in x.text for x in s.shapes if x.has_text_frame) for s in temporal)


@pytest.mark.parametrize("kind", FORMATS)
@pytest.mark.parametrize("scenario", ["missing", "gap", "duplicate"])
def test_unavailable_comparison_never_becomes_zero_or_comparable(monkeypatch, kind, scenario):
    db = _Db()
    _enable(monkeypatch, db)
    current, _ = save(db, 1, "2026", -40, 12)
    if scenario != "missing":
        save(db, 1, "2024" if scenario == "gap" else "2026", -110, 11)
    response = asyncio.run(get(kind, current))
    assert response.status_code == 200
    text = text_of(kind, response.content)
    assert "UNKNOWN temporel" in text if scenario != "duplicate" else "CONTRADICTION temporelle" in text
    assert "EBITDA - ecart arithmetique" not in text
    assert "Ecarts arithmetiques disponibles" not in text


@pytest.mark.parametrize("kind", FORMATS)
def test_foreign_scope_tampering_and_history_outage_refuse_export(monkeypatch, kind):
    db = _Db()
    _enable(monkeypatch, db)
    current, _ = save(db, 1, "2026", -40, 12)
    previous, _ = save(db, 1, "2025", -110, 11)
    foreign, _ = save(db, 3, "2025", 912345, 31, COMPANY_B)
    assert asyncio.run(get(kind, foreign)).status_code == 404
    row = next(r for r in db.tables["governed_analysis_envelopes"] if r["analysis_id"] == previous)
    original = copy.deepcopy(row["envelope_json"])
    row["envelope_json"]["source_facts"]["facts"][0]["value"] = 9999
    assert asyncio.run(get(kind, current)).status_code == 503
    row["envelope_json"] = original
    execute = _Query.execute
    def outage(query):
        if query.table == "governed_analysis_envelopes" and not any(k == "analysis_id" for k, _ in query.filters):
            raise RuntimeError("synthetic history unavailable")
        return execute(query)
    monkeypatch.setattr(_Query, "execute", outage)
    before = copy.deepcopy(db.tables)
    assert asyncio.run(get(kind, current)).status_code == 503
    assert db.tables == before


@pytest.mark.parametrize("kind", FORMATS)
def test_renderer_rejects_misbound_or_broadened_snapshot(kind):
    prior, current = synthetic_envelope("2025", -110), synthetic_envelope("2026", -40)
    comparison = compare_governed_envelopes(prior, current, previous_analysis_id="prior", current_analysis_id="current")
    for key, value in [("current_analysis_id", "foreign"), ("current_period", "2027"),
                       ("financial_comparability", "ESTABLISHED"), ("causal_interpretation", "Invented cause")]:
        forged = copy.deepcopy(comparison)
        forged[key] = value
        with pytest.raises(ValueError):
            FORMATS[kind](current, "current", temporal_comparison=forged)
    forged = copy.deepcopy(comparison)
    forged["changes"][0]["current_value"] = 123456
    with pytest.raises(ValueError):
        FORMATS[kind](current, "current", temporal_comparison=forged)


@pytest.mark.parametrize("kind", FORMATS)
def test_partial_comparison_keeps_unknowns(kind):
    prior = synthetic_envelope("2025", -110)
    current = synthetic_envelope("2026", -40)
    # Bounded comparator output: removing one prior metric must stay partial.
    prior = prior.model_copy(update={"source_facts": prior.source_facts.model_copy(
        update={"facts": tuple(f for f in prior.source_facts.facts if f.metric == "EBITDA")})})
    comparison = compare_governed_envelopes(prior, current, previous_analysis_id="prior", current_analysis_id="current")
    assert comparison["status"] == "PARTIALLY_COMPARABLE"
    text = text_of(kind, FORMATS[kind](current, "current", temporal_comparison=comparison))
    assert "Ecarts arithmetiques partiels" in text
    assert "UNKNOWN temporel" in text
    assert "WORKING_CAPITAL" in text
