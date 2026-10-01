"""One-shot live V41 injected rehearsal; default mode is local preparation only.

No OpenAI SDK, provider credential or external transport is imported here.
The only network target accepted is the fixed Integration Test Supabase URL.
Owner policy insert/disable actions are performed from separately generated SQL;
this runner merely observes their exact effects and never gains owner authority.
"""

from __future__ import annotations

import argparse
from copy import deepcopy
from datetime import datetime, timedelta, timezone
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import sys
import time
from types import SimpleNamespace
from typing import Any, Callable
from uuid import UUID

from fastapi import FastAPI, Header, HTTPException, Request
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader
from supabase import ClientOptions, create_client

from routers.governed_output import build_governed_output_router
from sandbox.heterogeneous_workbooks import _mock_response
from sandbox.v41_injected_rehearsal import (
    CHECKPOINT_HEAD,
    FAILED_ANALYSIS_ID,
    FAILED_POLICY_CONTRACT_BINDING_SHA256,
    FAILED_POLICY_EVIDENCE_SHA256,
    FAILED_POLICY_ID,
    FIXTURE_NAME,
    FIXTURE_SHA256,
    OWNER_ACTION_WINDOW_SECONDS,
    PROJECT_URL,
    SCOPE,
    file_sha256,
    freeze_manifest,
    read_manifest,
    write_new,
    specification_without_evidence,
)
from sandbox.v41_owner_handoff import OwnerPolicyHandoff
from services.bounded_producer_policy import verify_bounded_local_test_policy
from services.durable_generic_producer_admission import (
    DurableGenericAdmissionRefused,
    DurableGenericProducerAdmission,
    PreparedGenericExecution,
)
from services.generic_producer_candidate import (
    CONTRACT_BINDING,
    GENERIC_ADMISSION_CONTRACT_SHA256,
    InjectedOpenAIResponsesCandidate,
    PRODUCER_ID,
    PRODUCER_VERSION,
    TASK_ID,
    TASK_VERSION,
    freeze_generic_producer_request,
)
from services.governed_producer_adapter import GovernedProducerInvocationV2
from services.governed_producer_admission import (
    DURABLE_CONTRACT,
    LocalProducerAdmissionPreparation,
    LocalSyntheticProfile,
)
from services.producer_execution_contract import ExecutionBindingsV2
from services.v1_analysis_contract import UnderstandingResult


REPO = Path("C:/Users/ADMIN-FRED/Documents/Codex/Pepperyn-development")
RUNTIME = Path("C:/Users/ADMIN-FRED/Documents/Codex/Pepperyn-runtime")
ATTEMPT = RUNTIME / "v41-injected-4"
ENTRYPOINT = Path(__file__).resolve()
FIXTURE = REPO / "backend/tests/golden/fixtures" / FIXTURE_NAME
BASELINE_TABLES = (
    "analyses", "governed_analysis_envelopes", "governed_execution_receipts",
    "producer_policies_v2", "execution_admissions_v2", "execution_receipts_v2",
)
V41_TABLES = (
    "generic_producer_policies_v3", "generic_execution_admissions_v3",
    "generic_execution_receipts_v3",
)
AUTHORIZATION = "V41_INJECTED_DURABLE_SUCCESSOR_GO_10_EFFECTS_5_ROWS"


def require(value: Any, stage: str = "V41_REHEARSAL_REFUSED") -> None:
    if not value:
        raise ValueError(stage)


def digest(value: Any) -> str:
    raw = json.dumps(value, ensure_ascii=False, allow_nan=False,
                     sort_keys=True, separators=(",", ":")).encode("utf-8")
    return sha256(raw).hexdigest().upper()


def client(key: str):
    require(type(key) is str and bool(key.strip()), "V41_KEY_REFUSED")
    return create_client(PROJECT_URL, key, options=ClientOptions(
        auto_refresh_token=False, persist_session=False,
        postgrest_client_timeout=30, function_client_timeout=30,
    ))


def rows(db, table: str, **filters: str) -> list[dict]:
    query = db.from_(table).select("*")
    for key, value in filters.items():
        query = query.eq(key, value)
    data = query.limit(2001).execute().data
    require(isinstance(data, list) and len(data) < 2000 and
            all(isinstance(row, dict) for row in data), "V41_READ_REFUSED")
    return data


def snapshot(db) -> dict[str, dict[str, Any]]:
    result = {}
    for table in BASELINE_TABLES:
        observed = rows(db, table)
        result[table] = {"count": len(observed), "sha256": digest(sorted(observed, key=digest))}
    return result


def scope(db, company_name: str) -> dict[str, str]:
    companies = rows(db, "companies", name=company_name)
    require(len(companies) == 1, "V41_COMPANY_SCOPE_REFUSED")
    company_id = str(UUID(companies[0]["id"]))
    actor_id = str(UUID(companies[0]["admin_user_id"]))
    profiles = rows(db, "profiles", id=actor_id)
    entities = rows(db, "entities", company_id=company_id)
    require(len(profiles) == 1 and profiles[0].get("company_id") == company_id,
            "V41_PROFILE_SCOPE_REFUSED")
    require(len(entities) == 1 and entities[0].get("company_id") == company_id,
            "V41_ENTITY_SCOPE_REFUSED")
    entity_id = str(UUID(entities[0]["id"]))
    engagements = rows(db, "engagements", entity_id=entity_id)
    require(len(engagements) == 1 and engagements[0].get("entity_id") == entity_id,
            "V41_ENGAGEMENT_SCOPE_REFUSED")
    return {"actor_id": actor_id, "company_id": company_id, "entity_id": entity_id,
            "engagement_id": str(UUID(engagements[0]["id"]))}


class EffectBudget:
    def __init__(self, path: Path) -> None:
        self.path = path
        self.effects: list[dict[str, Any]] = []
        self.auth_logins = 0
        self._persist()

    def _persist(self) -> None:
        payload = {"ceiling": 10, "auth_ceiling": 1, "effects": self.effects,
                   "auth_logins": self.auth_logins}
        temp = self.path.with_suffix(".tmp")
        temp.write_text(json.dumps(payload, sort_keys=True, separators=(",", ":")) + "\n",
                        encoding="utf-8", newline="\n")
        os.replace(temp, self.path)

    def start(self, name: str, *, auth: bool = False) -> int:
        require(len(self.effects) < 10, "V41_EFFECT_CEILING_REFUSED")
        if auth:
            require(self.auth_logins == 0, "V41_SECOND_AUTH_REFUSED")
            self.auth_logins = 1
        self.effects.append({"number": len(self.effects) + 1, "name": name, "outcome": "STARTED"})
        self._persist()
        return len(self.effects) - 1

    def finish(self, index: int, outcome: str) -> None:
        require(self.effects[index]["outcome"] == "STARTED", "V41_EFFECT_JOURNAL_REFUSED")
        self.effects[index]["outcome"] = outcome
        self._persist()

    def run(self, name: str, action: Callable[[], Any], *, auth: bool = False,
            refusal: bool = False) -> Any:
        index = self.start(name, auth=auth)
        try:
            value = action()
        except Exception:
            if refusal:
                self.finish(index, "REFUSED_AS_REQUIRED")
                return None
            self.finish(index, "FAILED_OR_UNCERTAIN")
            raise
        if refusal:
            self.finish(index, "UNEXPECTED_ACCEPTANCE")
            raise ValueError("V41_EXPECTED_REFUSAL_ACCEPTED")
        self.finish(index, "ACKNOWLEDGED")
        return value

    def owner_observed(self, name: str) -> None:
        index = self.start(name)
        self.finish(index, "OBSERVED_EXACT_EFFECT")


class RecordingDb:
    def __init__(self, db) -> None:
        self.db = db
        self.rpc_params: dict[str, dict[str, Any]] = {}

    @property
    def auth(self):
        # Preserve backend token verification through the recording adapter.
        # No login, cached principal or producer-provided identity is introduced.
        return self.db.auth

    def from_(self, table):
        return self.db.from_(table)

    def rpc(self, name, params):
        self.rpc_params[name] = deepcopy(params)
        return self.db.rpc(name, params)


def precontrol_attestation(manifest: dict[str, Any]) -> dict[str, Any]:
    path = ATTEMPT / "precontrol-ready.json"
    require(path.is_file(), "V41_PRECONTROL_ATTESTATION_ABSENT")
    report = json.loads(path.read_text(encoding="utf-8"))
    expected_keys = {
        "schema_version", "project_url", "sql_sha256", "structural_status",
        "historical_status", "frozen_scope_status", "historical_snapshot_sha256",
        "observed_at", "write_performed", "auth_performed",
    }
    bounded = report.get("schema_version") == "v41-injected-successor-precontrol-attestation-2"
    if bounded:
        expected_keys = (expected_keys - {"observed_at"}) | {"freshness_evidence"}
    require(set(report) == expected_keys, "V41_PRECONTROL_ATTESTATION_SHAPE_REFUSED")
    require((bounded or report["schema_version"] == "v41-injected-successor-precontrol-attestation-1") and
            report["project_url"] == PROJECT_URL and
            report["sql_sha256"] == file_sha256(ATTEMPT / "precontrol.sql") and
            report["structural_status"] == "PPR067_HARDENING_STRUCTURAL_POSTFLIGHT_PASS" and
            report["historical_status"] == "V41_HISTORICAL_BASELINE_PASS" and
            report["frozen_scope_status"] == "V41_INJECTED_SUCCESSOR_FROZEN_SCOPE_PRECONTROL_PASS" and
            isinstance(report["historical_snapshot_sha256"], str) and
            len(report["historical_snapshot_sha256"]) == 64 and
            report["write_performed"] is False and report["auth_performed"] is False,
            "V41_PRECONTROL_ATTESTATION_REFUSED")
    if bounded:
        validate_conservative_freshness(report, manifest, datetime.now(timezone.utc))
    else:
        stamp = datetime.fromisoformat(report["observed_at"])
        require(stamp.utcoffset() is not None and
                0 <= (datetime.now(timezone.utc) - stamp).total_seconds()
                <= OWNER_ACTION_WINDOW_SECONDS,
                "V41_PRECONTROL_ATTESTATION_EXPIRED")
    require(manifest["checkpoint_head"] == CHECKPOINT_HEAD, "V41_CHECKPOINT_REFUSED")
    return report


def validate_conservative_freshness(report, manifest, now):
    """Founder-qualified attempt-4 bound; never an execution/observation time."""
    evidence = report.get("freshness_evidence")
    expected = {
        "kind": "PROVEN_CONSERVATIVE_LOWER_BOUND",
        "lower_bound_utc": "2026-10-01T07:24:01Z",
        "qualification": "NOT_EXECUTION_OR_OBSERVATION_TIME",
        "authority": "FOUNDER_ATTEMPT4_CONSERVATIVE_BOUND_2026_10_01",
        "provenance": "FROZEN_IDENTITIES_PRESENT_IN_ACCEPTED_PRECONTROL_RESULT",
        "manifest_sha256": "750A16888ADC70D2EFC4F3AF7199CC821458E1C75715A76999CAACD451D890CA",
        "result_file": "accepted-precontrol-result.json",
        "result_sha256": "7205B6612A16746E768AFCC596D5D0B193DF547CAC08E2748A301F8ABE0C3853",
    }
    require(evidence == expected and "observed_at" not in report,
            "V41_CONSERVATIVE_EVIDENCE_REFUSED")
    # Revalidate the immutable manifest rather than accepting caller assertions.
    require(read_manifest(ATTEMPT) == manifest and
            manifest["manifest_sha256"] == expected["manifest_sha256"],
            "V41_CONSERVATIVE_MANIFEST_REFUSED")
    path = ATTEMPT / expected["result_file"]
    require(file_sha256(path) == evidence["result_sha256"],
            "V41_CONSERVATIVE_RESULT_HASH_REFUSED")
    result = json.loads(path.read_text(encoding="utf-8"))
    require(result.get("check_version") == "v41-successor-single-result-precontrol-2" and
            result.get("status") == "V41_SUCCESSOR_PRECONTROL_CHECKS_PASS" and
            result.get("auth_performed") is False and result.get("write_performed") is False and
            result.get("historical_comparison_required") is True and
            digest(result["historical"]) == report["historical_snapshot_sha256"],
            "V41_CONSERVATIVE_RESULT_REFUSED")
    for field, component in (("structural_status", "structural"),
                             ("historical_status", "historical"),
                             ("frozen_scope_status", "frozen_scope")):
        require(result[component]["status"] == report[field], "V41_CONSERVATIVE_STATUS_REFUSED")
    require(all(result["frozen_scope"].get(k) == v for k, v in manifest["identities"].items()),
            "V41_CONSERVATIVE_IDENTITIES_REFUSED")
    bound = datetime.fromisoformat(expected["lower_bound_utc"].replace("Z", "+00:00"))
    created = datetime.fromisoformat(manifest["created_at"])
    deadline = datetime.fromisoformat(manifest["owner_action_deadline"])
    require(now.utcoffset() is not None and bound <= created <= now < deadline and
            0 <= (now - bound).total_seconds() < 86400,
            "V41_CONSERVATIVE_FRESHNESS_REFUSED")


def failed_policy_exact(row: dict[str, Any]) -> bool:
    specification = row.get("specification")
    return (
        row.get("id") == FAILED_POLICY_ID
        and row.get("enabled") is False
        and row.get("contract_binding_sha256")
        == FAILED_POLICY_CONTRACT_BINDING_SHA256
        and isinstance(specification, dict)
        and specification.get("policy_evidence_sha256")
        == FAILED_POLICY_EVIDENCE_SHA256
    )


def prepolicy_remote(db, manifest: dict[str, Any]) -> tuple[dict, dict]:
    policies = rows(db, "generic_producer_policies_v3")
    require(len(policies) == 1 and failed_policy_exact(policies[0]),
            "V41_FAILED_POLICY_HISTORY_REFUSED")
    require(rows(db, "generic_execution_admissions_v3") == [] and
            rows(db, "generic_execution_receipts_v3") == [],
            "V41_SUCCESSOR_REGISTRIES_REFUSED")
    own = scope(db, "Pepperyn A24 Isolation Synthetic 1")
    other = scope(db, "Pepperyn A24 Isolation Synthetic 2")
    require(own == manifest["scope"] and own != other, "V41_SCOPE_DRIFT_REFUSED")
    ids = manifest["identities"]
    require(rows(db, "generic_producer_policies_v3", id=ids["policy_id"]) == [] and
            rows(db, "generic_execution_admissions_v3", execution_id=ids["execution_id"]) == [] and
            rows(db, "analyses", id=ids["analysis_id"]) == [] and
            rows(db, "governed_analysis_envelopes", analysis_id=ids["analysis_id"]) == [] and
            rows(db, "generic_execution_receipts_v3", analysis_id=ids["analysis_id"]) == [],
            "V41_FROZEN_IDENTITY_PRESENT")
    require(rows(db, "analyses", id=FAILED_ANALYSIS_ID) == [] and
            rows(db, "governed_analysis_envelopes", analysis_id=FAILED_ANALYSIS_ID) == [],
            "V41_FAILED_ATTEMPT_APPLICATION_ROW_REFUSED")
    return own, other


def postpolicy_remote(db, manifest: dict[str, Any], baseline: dict[str, Any], policies) -> dict:
    policy_id = manifest["identities"]["policy_id"]
    failed = [row for row in policies if row.get("id") == FAILED_POLICY_ID]
    successor = [row for row in policies if row.get("id") == policy_id]
    require(len(policies) == 2 and len(failed) == 1 and
            failed_policy_exact(failed[0]) and len(successor) == 1 and
            successor[0].get("enabled") is True,
            "V41_SUCCESSOR_POLICY_SET_REFUSED")
    require(rows(db, "generic_execution_admissions_v3") == [] and
            rows(db, "generic_execution_receipts_v3") == [] and
            snapshot(db) == baseline,
            "V41_SUCCESSOR_POSTPOLICY_STATE_REFUSED")
    return successor[0]


def prepare_execution(db, token: str, manifest: dict[str, Any]) -> tuple[PreparedGenericExecution, bytes]:
    ids, own = manifest["identities"], manifest["scope"]
    raw = FIXTURE.read_bytes()
    require(sha256(raw).hexdigest().upper() == FIXTURE_SHA256, "V41_FIXTURE_REFUSED")
    profile = LocalSyntheticProfile(
        company_id=own["company_id"], entity_id=own["entity_id"],
        engagement_id=own["engagement_id"], producer_id=PRODUCER_ID,
        producer_version=PRODUCER_VERSION, task_id=TASK_ID, task_version=TASK_VERSION,
        source_sha256=FIXTURE_SHA256, filename=FIXTURE_NAME,
    )
    identities = iter((UUID(ids["analysis_id"]), UUID(ids["request_id"]), UUID(ids["execution_id"])))
    prep = LocalProducerAdmissionPreparation(
        db, profile=profile, contract_policy=DURABLE_CONTRACT,
        identity_source=lambda: tuple(next(identities) for _ in range(3)),
    ).prepare(authorization="Bearer " + token, entity_id=own["entity_id"],
              engagement_id=own["engagement_id"], raw=raw, filename=FIXTURE_NAME)
    facts = UnderstandingResult.model_validate_json(prep.input_json)
    bindings = prep.bindings.model_copy(update={
        "producer_id": PRODUCER_ID, "producer_version": PRODUCER_VERSION,
        "task_id": TASK_ID, "task_version": TASK_VERSION,
        "admission_contract_sha256": GENERIC_ADMISSION_CONTRACT_SHA256,
        "raw_source_sha256": FIXTURE_SHA256,
    })
    invocation = GovernedProducerInvocationV2(
        task_id=TASK_ID, task_version=TASK_VERSION,
        invocation_nonce=bindings.request_id.hex.upper(), source_facts=facts,
    )
    frozen = freeze_generic_producer_request(invocation)
    bindings = bindings.model_copy(update={"producer_input_sha256": frozen.request_sha256})
    return PreparedGenericExecution(
        policy_id=UUID(ids["policy_id"]), bindings=bindings,
        source_facts=facts, frozen_request=frozen, filename=FIXTURE_NAME,
    ), raw


def rpc_refusal(db, prepared: PreparedGenericExecution, bindings: ExecutionBindingsV2) -> None:
    db.rpc("reserve_generic_execution_v3", {
        "p_policy": str(prepared.policy_id),
        "p_bindings": bindings.model_dump(mode="json"),
        "p_contract_binding": CONTRACT_BINDING.model_dump(mode="json"),
        "p_source_facts": prepared.source_facts.model_dump(mode="json"),
        "p_projection_text": prepared.frozen_request.canonical_request,
        "p_filename": prepared.filename, "p_ttl": 300,
    }).execute()


def wait_policy(db, manifest: dict[str, Any], *, enabled: bool) -> dict:
    if enabled:
        deadline = datetime.fromisoformat(manifest["owner_action_deadline"])
        require(deadline.utcoffset() is not None, "V41_OWNER_ACTION_DEADLINE_REFUSED")
    else:
        deadline = datetime.now(timezone.utc) + timedelta(seconds=1200)
    action = "insert" if enabled else "disable"
    artifact = ATTEMPT / f"policy-{action}.sql"
    expected = {
        "schema_version": "v41-owner-local-ack-1",
        "action": action,
        "manifest_sha256": manifest["manifest_sha256"],
        "policy_id": manifest["identities"]["policy_id"],
        "sql_sha256": file_sha256(artifact),
    }
    write_new(ATTEMPT / f"owner-{action}-handoff-started.json", {
        "ack": expected, "deadline_utc": deadline.isoformat(),
    })
    original = deepcopy(manifest)

    def freshness():
        require(read_manifest(ATTEMPT) == original, "V41_OWNER_MANIFEST_CHANGED")
        require(file_sha256(artifact) == expected["sql_sha256"], "V41_OWNER_SQL_CHANGED")
        if enabled:
            precontrol_attestation(original)

    def acknowledge(check_time):
        path = ATTEMPT / f"owner-{action}-ack.json"
        while not path.exists():
            check_time()
            time.sleep(1)  # Local filesystem only during human intervention.
        return json.loads(path.read_text(encoding="utf-8"))

    handoff = OwnerPolicyHandoff(expected, deadline=deadline, freshness=freshness,
                                 clock=lambda: datetime.now(timezone.utc))
    return handoff.observe(acknowledge,
        lambda: rows(db, "generic_producer_policies_v3"),
        lambda policies: validate_observed_owner_policies(policies, original, enabled=enabled))


def validate_observed_owner_policies(policies, manifest, *, enabled):
    require(type(policies) is list and len(policies) == 2, "V41_OWNER_POLICY_SET_REFUSED")
    historical = [p for p in policies if p.get("id") == FAILED_POLICY_ID]
    current = [p for p in policies if p.get("id") == manifest["identities"]["policy_id"]]
    require(len(historical) == 1 and failed_policy_exact(historical[0]) and
            len(current) == 1, "V41_OWNER_POLICY_ID_REFUSED")
    policy = current[0]
    spec = policy.get("specification")
    require(type(spec) is dict and
            set(spec) == set(specification_without_evidence(manifest)) | {"policy_evidence_sha256"},
            "V41_OWNER_SPEC_REFUSED")
    require({k: v for k, v in spec.items() if k != "policy_evidence_sha256"}
            == specification_without_evidence(manifest), "V41_OWNER_SCOPE_REFUSED")
    require(policy.get("enabled") is enabled and
            policy.get("contract_binding") == CONTRACT_BINDING.model_dump(mode="json") and
            policy.get("contract_binding_sha256") == GENERIC_ADMISSION_CONTRACT_SHA256,
            "V41_OWNER_CONTRACT_REFUSED")
    from services.bounded_producer_policy import BoundedLocalTestPolicyV1
    BoundedLocalTestPolicyV1.model_validate(spec)


def exact_v41_rows(db, manifest: dict[str, Any], *, state: str, enabled: bool) -> dict[str, list[dict]]:
    ids = manifest["identities"]
    result = {
        "policy": rows(db, "generic_producer_policies_v3", id=ids["policy_id"]),
        "admission": rows(db, "generic_execution_admissions_v3", execution_id=ids["execution_id"]),
        "analysis": rows(db, "analyses", id=ids["analysis_id"]),
        "envelope": rows(db, "governed_analysis_envelopes", analysis_id=ids["analysis_id"]),
        "receipt": rows(db, "generic_execution_receipts_v3", analysis_id=ids["analysis_id"]),
    }
    require(all(len(value) == 1 for value in result.values()), "V41_EXACT_ROWS_REFUSED")
    require(result["policy"][0].get("enabled") is enabled and
            result["admission"][0].get("state") == state,
            "V41_TERMINAL_STATE_REFUSED")
    return result


def independent_recovery(packet: dict[str, Any]) -> dict[str, Any]:
    client(packet["anon"])
    service = client(packet["service"])
    token = packet["token"]
    user = service.auth.get_user(token).user
    require(user is not None and str(UUID(user.id)) == SCOPE["actor_id"], "V41_RECOVERY_ACTOR_REFUSED")
    profile = rows(service, "profiles", id=str(UUID(user.id)))
    require(len(profile) == 1 and profile[0].get("company_id") == SCOPE["company_id"],
            "V41_RECOVERY_COMPANY_REFUSED")
    analysis_id = packet["analysis_id"]

    def authenticated_company(authorization: str | None = Header(default=None)):
        if authorization != "Bearer " + token:
            raise HTTPException(404, "Ressource introuvable")
        return SCOPE["company_id"]

    def exact_get(request: Request):
        if request.method != "GET" or request.path_params.get("analysis_id") != analysis_id:
            raise HTTPException(503, "Lecture fermee")

    app = FastAPI()
    app.include_router(build_governed_output_router(
        authenticated_company=authenticated_company,
        require_admission=exact_get,
        database=lambda: service,
    ))
    transport = TestClient(app)
    headers = {"Authorization": "Bearer " + token}
    reread = transport.get(f"/api/governed/analyses/{analysis_id}", headers=headers)
    require(reread.status_code == 200, "V41_OWNER_REREAD_REFUSED")
    body = reread.json()
    provenance = body.get("execution_provenance")
    require(body.get("analysis_id") == analysis_id and
            provenance.get("receipt_version") == "V41" and
            provenance["receipt"].get("provider_execution_attested") is False and
            provenance["receipt"].get("producer_admission_status") == "UNADMITTED",
            "V41_OWNER_PROVENANCE_REFUSED")
    exports = {}
    for extension in ("xlsx", "pdf", "pptx"):
        response = transport.get(
            f"/api/governed/analyses/{analysis_id}/export.{extension}", headers=headers)
        require(response.status_code == 200 and len(response.content) > 500,
                "V41_EXPORT_REFUSED")
        if extension == "xlsx":
            workbook = load_workbook(BytesIO(response.content), data_only=True)
            text = "\n".join(
                str(value) for sheet in workbook for row in sheet.iter_rows(values_only=True)
                for value in row if value is not None
            )
        elif extension == "pdf":
            text = "\n".join(
                page.extract_text() or "" for page in PdfReader(BytesIO(response.content)).pages
            )
        else:
            text = "\n".join(
                shape.text for slide in Presentation(BytesIO(response.content)).slides
                for shape in slide.shapes if shape.has_text_frame
            )
        required = (
            "V41 / governed-generic-producer-receipt-3",
            "Reponse injectee locale",
            "aucune execution OpenAI attestee",
            "Non admis globalement",
            "Reseau externe",
            "Aucun",
        )
        require(all(value in text for value in required), "V41_EXPORT_PROVENANCE_REFUSED")
        exports[extension] = {"bytes": len(response.content),
                              "sha256": sha256(response.content).hexdigest().upper(),
                              "provenance_verified": True}
    fixture = ATTEMPT / "ui-provenance.json"
    write_new(fixture, provenance)
    return {"status": "V41_INDEPENDENT_RECOVERY_PASS", "analysis_id": analysis_id,
            "receipt_version": "V41", "provider_execution_attested": False,
            "producer_admission_status": "UNADMITTED", "exports": exports,
            "ui_fixture": str(fixture)}


def recover_main() -> int:
    packet = json.load(sys.stdin)
    require(set(packet) == {"anon", "service", "token", "analysis_id"},
            "V41_RECOVERY_PACKET_REFUSED")
    result = independent_recovery(packet)
    print(json.dumps(result, sort_keys=True))
    return 0


def execute(packet: dict[str, Any]) -> int:
    stage = "START"
    budget = None
    service = None
    manifest = read_manifest(ATTEMPT)
    policy_observed = False
    policy_disabled = False
    baseline = None
    try:
        require(set(packet) == {"anon", "service", "bundle", "authorization"},
                "V41_PACKET_REFUSED")
        require(packet["authorization"] == AUTHORIZATION, "V41_AUTHORIZATION_REFUSED")
        bundle = packet["bundle"]
        require(bundle.get("project_url") == PROJECT_URL and
                bundle.get("purpose") == "A24_TECHNICAL_ISOLATION_ONLY" and
                len(bundle.get("accounts", [])) == 2 and
                bundle["accounts"][0].get("email") ==
                "pepperyn-isolation-a24-a@pepperyn-test.invalid",
                "V41_DPAPI_BUNDLE_REFUSED")
        require(not any((ATTEMPT / name).exists() for name in (
            "refused.json", "result.json", "effects.json", "baseline-before.json")),
            "V41_ATTEMPT_ALREADY_STARTED")
        precontrol_attestation(manifest)
        service = client(packet["service"])
        stage = "FRESH_PREPOLICY_REMOTE"
        own, foreign = prepolicy_remote(service, manifest)
        baseline = snapshot(service)
        write_new(ATTEMPT / "baseline-before.json", baseline)
        budget = EffectBudget(ATTEMPT / "effects.json")

        stage = "POLICY_INSERT"
        print("V41_INJECTED_OWNER_ACTION_REQUIRED: POLICY_INSERT", flush=True)
        observed_policies = wait_policy(service, manifest, enabled=True)
        policy_observed = True
        budget.owner_observed("POLICY_INSERT")
        policy = postpolicy_remote(service, manifest, baseline, observed_policies)

        # Waiting for owner action cannot renew the precontrol's freshness.
        precontrol_attestation(manifest)

        stage = "SINGLE_AUTH"
        anon = client(packet["anon"])
        auth = budget.run("AUTH_LOGIN", lambda: anon.auth.sign_in_with_password({
            "email": bundle["accounts"][0]["email"],
            "password": bundle["accounts"][0]["password"],
        }), auth=True)
        token = auth.session.access_token
        require(auth.user is not None and str(UUID(auth.user.id)) == own["actor_id"],
                "V41_AUTH_ACTOR_REFUSED")

        stage = "COMPOSITION"
        recorded = RecordingDb(service)
        prepared, raw = prepare_execution(recorded, token, manifest)
        verify_bounded_local_test_policy(
            policy, bindings=prepared.bindings, filename=prepared.filename,
            require_enabled=True,
        )
        principal = lambda _: SimpleNamespace(
            principal_id=own["actor_id"], company_id=own["company_id"])
        admission = DurableGenericProducerAdmission(recorded, principal_resolver=principal)

        stage = "FOREIGN_SCOPE_REFUSAL"
        foreign_bindings = prepared.bindings.model_copy(update={
            "company_id": foreign["company_id"], "entity_id": foreign["entity_id"],
            "engagement_id": foreign["engagement_id"],
        })
        budget.run("FOREIGN_SCOPE_RESERVE", lambda: rpc_refusal(
            service, prepared, foreign_bindings), refusal=True)
        require(rows(service, "generic_execution_admissions_v3") == [],
                "V41_FOREIGN_REFUSAL_EFFECT_REFUSED")

        stage = "SUBSTITUTED_REQUEST_REFUSAL"
        substituted = prepared.bindings.model_copy(update={"producer_input_sha256": "A" * 64})
        budget.run("SUBSTITUTED_REQUEST_RESERVE", lambda: rpc_refusal(
            service, prepared, substituted), refusal=True)
        require(rows(service, "generic_execution_admissions_v3") == [],
                "V41_SUBSTITUTION_REFUSAL_EFFECT_REFUSED")

        stage = "EXACT_RESERVE"
        reservation = budget.run("EXACT_RESERVE", lambda: admission.reserve(
            prepared, authorization="Bearer " + token, raw_source=raw))
        stage = "EXACT_CLAIM"
        claimed = budget.run("EXACT_CLAIM", lambda: admission.claim(
            reservation, authorization="Bearer " + token))
        stage = "INJECTED_COMPLETE"
        producer = InjectedOpenAIResponsesCandidate(lambda request: _mock_response(
            prepared.source_facts, prepared.bindings.request_id.hex.upper()))
        result = budget.run("INJECTED_COMPLETE", lambda: admission.execute_local_injected(
            claimed, authorization="Bearer " + token, producer=producer))
        require(result.get("provider_execution_attested") is False and
                result.get("external_provider_used") is False,
                "V41_INJECTED_RESULT_REFUSED")
        exact_v41_rows(service, manifest, state="COMPLETE", enabled=True)

        stage = "REPLAY_REFUSAL"
        complete_params = deepcopy(recorded.rpc_params["complete_generic_execution_v3"])
        budget.run("COMPLETION_REPLAY", lambda: service.rpc(
            "complete_generic_execution_v3", complete_params).execute(), refusal=True)
        exact_v41_rows(service, manifest, state="COMPLETE", enabled=True)

        stage = "POLICY_DISABLE"
        print("V41_INJECTED_OWNER_ACTION_REQUIRED: POLICY_DISABLE", flush=True)
        wait_policy(service, manifest, enabled=False)
        policy_disabled = True
        budget.owner_observed("POLICY_DISABLE")

        stage = "POST_DISABLE_REFUSAL"
        budget.run("POST_DISABLE_RESERVE", lambda: rpc_refusal(
            service, prepared, prepared.bindings), refusal=True)
        exact = exact_v41_rows(service, manifest, state="COMPLETE", enabled=False)

        stage = "INDEPENDENT_RECOVERY"
        recovery_packet = {"anon": packet["anon"], "service": packet["service"],
                           "token": token, "analysis_id": manifest["identities"]["analysis_id"]}
        recovered = subprocess.run(
            [sys.executable, "-B", str(ENTRYPOINT), "--recover"],
            input=json.dumps(recovery_packet), text=True, capture_output=True,
            timeout=180, check=False, env={**os.environ, "PYTHONPATH": str(REPO / "backend")},
        )
        require(recovered.returncode == 0, "V41_INDEPENDENT_PROCESS_REFUSED")
        recovery = json.loads(recovered.stdout.strip().splitlines()[-1])
        require(recovery.get("status") == "V41_INDEPENDENT_RECOVERY_PASS",
                "V41_INDEPENDENT_RECOVERY_REFUSED")

        stage = "UI_COMPONENT"
        ui_env = {**os.environ,
                  "PEPPERYN_V41_LIVE_UI_FIXTURE": recovery["ui_fixture"]}
        ui = subprocess.run(
            ["npm.cmd", "test", "--", "--runInBand",
             "components/chat/__tests__/ExecutionProvenance.rehearsal.test.tsx"],
            cwd=REPO / "frontend", env=ui_env, capture_output=True, text=True,
            timeout=180, check=False,
        )
        require(ui.returncode == 0, "V41_UI_COMPONENT_REFUSED")

        stage = "POSTCONTROL"
        after = snapshot(service)
        require(after == baseline, "V41_HISTORICAL_BASELINE_CHANGED")
        require(len(budget.effects) == 10 and budget.auth_logins == 1 and
                [entry["name"] for entry in budget.effects] == [
                    "POLICY_INSERT", "AUTH_LOGIN", "FOREIGN_SCOPE_RESERVE",
                    "SUBSTITUTED_REQUEST_RESERVE", "EXACT_RESERVE", "EXACT_CLAIM",
                    "INJECTED_COMPLETE", "COMPLETION_REPLAY", "POLICY_DISABLE",
                    "POST_DISABLE_RESERVE",
                ], "V41_EFFECT_SEQUENCE_REFUSED")
        require(sum(len(value) for value in exact.values()) == 5,
                "V41_DURABLE_ROW_CEILING_REFUSED")
        output = {
            "status": "BOUNDED_V41_INJECTED_DURABLE_SUCCESSOR_REHEARSAL_PASS",
            "analysis_id": manifest["identities"]["analysis_id"],
            "new_durable_rows": 5, "effect_attempts": 10, "auth_logins": 1,
            "adversarial_refusals": 3, "policy_enabled": False,
            "admission_state": "COMPLETE", "execution_evidence_status": "ADMITTED_EXECUTION",
            "execution_admission_only": True, "producer_global_status": "UNADMITTED",
            "transport_mode": "INJECTED_LOCAL_ONLY",
            "provider_execution_attested": False, "independent_recovery": recovery,
            "ui_component_verified": True, "historical_baseline_unchanged": True,
            "external_provider_used": False, "real_data_used": False,
            "b1_closed": False, "production_proven": False,
        }
        write_new(ATTEMPT / "result.json", output)
        print(json.dumps(output, sort_keys=True), flush=True)
        return 0
    except BaseException as error:
        if policy_observed and not policy_disabled and service is not None:
            print("V41_INJECTED_OWNER_ACTION_REQUIRED: SAFETY_DISABLE", flush=True)
            try:
                wait_policy(service, manifest, enabled=False)
                policy_disabled = True
                if budget is not None:
                    budget.owner_observed("SAFETY_POLICY_DISABLE")
            except Exception:
                pass
        failure = {
            "status": "V41_INJECTED_REHEARSAL_REFUSED", "stage": stage,
            "policy_observed": policy_observed, "policy_disabled": policy_disabled,
            "automatic_retry_permitted": False, "error_type": type(error).__name__,
            "observation_limit": "NOT_PROOF_OF_REMOTE_ABSENCE",
        }
        if not (ATTEMPT / "refused.json").exists():
            write_new(ATTEMPT / "refused.json", failure)
        print(json.dumps(failure, sort_keys=True), flush=True)
        return 1
    finally:
        packet.clear()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--prepare", action="store_true")
    parser.add_argument("--execute", action="store_true")
    parser.add_argument("--recover", action="store_true")
    args = parser.parse_args()
    require(sum((args.prepare, args.execute, args.recover)) == 1, "V41_MODE_REFUSED")
    if args.prepare:
        manifest = freeze_manifest(REPO, ATTEMPT)
        print(json.dumps({
            "status": "V41_INJECTED_PRECONTROL_READY", "project": PROJECT_URL,
            "analysis_id": manifest["identities"]["analysis_id"],
            "manifest_sha256": manifest["manifest_sha256"],
            "precontrol_sql_sha256": file_sha256(ATTEMPT / "precontrol.sql"),
            "auth_performed": False, "write_performed": False,
        }, sort_keys=True))
        return 0
    if args.recover:
        return recover_main()
    packet = json.load(sys.stdin)
    return execute(packet)


if __name__ == "__main__":
    raise SystemExit(main())
