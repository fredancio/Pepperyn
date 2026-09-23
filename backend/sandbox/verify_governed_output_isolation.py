"""A29: isolated ASGI GET routes, real sessions/persistence; not port-8000 proof.

One designated company per sequential phase. No admission guard replacement.
Database adapter denies mutations. Exports stay in memory; no business writes.
"""
import json
import os
from io import BytesIO
from unittest.mock import patch

from sandbox.inspect_populated_history_scope import inspect
from sandbox.seed_isolation_history import fixture_id
from sandbox.verify_feedback_privileges import URL


class ReadOnly:
    """Explicit allowlist, including auth.get_user but never mutation/RPC."""
    METHODS = frozenset(('table', 'from_', 'select', 'eq', 'is_', 'limit', 'order', 'in_', 'get_user'))

    def __init__(self, target): self._target = target
    def __getattr__(self, name):
        if name == 'auth': return ReadOnly(self._target.auth)
        if name == 'execute': return self._target.execute
        if name not in self.METHODS: raise RuntimeError('A29_READ_ONLY_REFUSED')
        def call(*args, **kwargs):
            result = getattr(self._target, name)(*args, **kwargs)
            return result if name == 'get_user' else ReadOnly(result)
        return call


def output_text(content, kind):
    if kind == 'xlsx':
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(content), read_only=True)
        try:
            return '\n'.join(str(c.value) for s in book for row in s for c in row if c.value is not None)
        finally: book.close()
    if kind == 'pptx':
        from pptx import Presentation
        return '\n'.join(s.text for slide in Presentation(BytesIO(content)).slides
                         for s in slide.shapes if s.has_text_frame)
    from pypdf import PdfReader
    return '\n'.join(p.extract_text() or '' for p in PdfReader(BytesIO(content)).pages)


def verify(bundle, anon, service, factory, *, live_transport=False):
    stage = 'INPUT'
    positives = denials = 0
    try:
        def require(value):
            if not value: raise ValueError('REFUSED')
        require(os.environ.get('ENVIRONMENT') == 'development')
        scopes = []
        stage = 'EXACT_SCOPES'
        require(inspect(bundle, service, anon, factory, verified_scopes=scopes)['status'] == 'A26_SCOPES_INSPECTED')
        require(len(scopes) == 2)
        db = ReadOnly(factory(URL, service))
        ids = [fixture_id(s) for s in scopes]
        stage = 'A26_MARKERS'
        for scope, aid in zip(scopes, ids):
            rows = db.table('analyses').select('id,company_id,entity_id,contexte_utilisateur').eq('id', aid).eq('company_id', scope[1]).limit(2).execute().data
            require(len(rows) == 1 and rows[0]['entity_id'] == scope[2]
                    and rows[0]['contexte_utilisateur'] == 'A26_TECHNICAL_SYNTHETIC_ISOLATION_FIXTURE_V1')
        tokens = []
        for account, scope in zip(bundle['accounts'], scopes):
            client = factory(URL, anon)
            session = client.auth.sign_in_with_password(account).session
            require(str(client.auth.get_user(session.access_token).user.id) == scope[0])
            tokens.append(session.access_token)
        import main
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from sandbox.v1_router import router
        app = FastAPI()
        allowed = {'/api/v1/governed-analyses/{analysis_id}' + suffix
                   for suffix in ('', '/export.xlsx', '/export.pdf', '/export.pptx')}
        app.router.routes = [r for r in router.routes if r.path in allowed and r.methods == {'GET'}]
        require(len(app.router.routes) == 4)
        # Only service transport is constrained. Authentication and ownership code
        # are unchanged and consult the real Supabase project using real sessions.
        from sandbox.output_live_transport import live_client
        transport = live_client(main.app, ids) if live_transport else TestClient(app)
        beta_env = {'PEPPERYN_PRIVATE_BETA': '1', 'PEPPERYN_BETA_USER_IDS': ','.join(s[0] for s in scopes)} if live_transport else {}
        with patch.object(main, 'get_supabase_service', lambda: db), patch.dict(os.environ, beta_env), transport as http:
            for i, scope in enumerate(scopes):
                with patch.dict(os.environ, {'PEPPERYN_SYNTHETIC_V1_COMPANY_ID': scope[1]}):
                    for kind in ('analysis', 'xlsx', 'pdf', 'pptx'):
                        stage = f'USER_{i+1}_{kind.upper()}'
                        suffix = '' if kind == 'analysis' else '/export.' + kind
                        path = '/api/v1/governed-analyses/' + ids[i] + suffix
                        own = {'Authorization': 'Bearer ' + tokens[i]}
                        response = http.get(path, headers=own)
                        require(response.status_code == 200)
                        text = json.dumps(response.json()) if kind == 'analysis' else output_text(response.content, kind)
                        require(ids[i] in text and ids[1-i] not in text)
                        positives += 1
                        # Same admitted identity requests foreign report: ownership,
                        # not designated-company denial, must reject the reference.
                        foreign = http.get('/api/v1/governed-analyses/' + ids[1-i] + suffix, headers=own)
                        require(foreign.status_code == 404 and foreign.json() == {'detail': 'Analyse introuvable'})
                        denials += 1
                        other = http.get(path, headers={'Authorization': 'Bearer ' + tokens[1-i]})
                        require(other.status_code == 404 and other.json() == {'detail': 'Ressource introuvable'})
                        denials += 1
                        require(http.get(path).status_code == 401)
                        denials += 1
        return dict(status='BOUNDED_A30_UVICORN_OUTPUT_ISOLATION_PASS' if live_transport else 'BOUNDED_A29_ASGI_OUTPUT_ISOLATION_PASS', positive_outputs=positives,
                    denials=denials, business_write_performed=False,
                    proof_scope='REAL_SESSIONS_A26_TEMPORARY_LOOPBACK_UVICORN_READ_SURFACE' if live_transport else 'REAL_USER_SESSIONS_PERSISTED_A26_ISOLATED_ASGI_GET_ONLY',
                    live_server_transport_proven=live_transport, global_isolation_proven=False,
                    write_isolation_proven=False, production_proof=False,
                    external_provider_used=False, real_data_used=False)
    except Exception:
        return dict(status='REFUSED', stage=stage, positive_outputs=positives, denials=denials,
                    business_write_performed=False, automatic_retry_permitted=False)


if __name__ == '__main__':
    import logging
    import warnings
    logging.disable(logging.CRITICAL)
    warnings.filterwarnings('ignore')
    try:
        from supabase import create_client
        from contextlib import redirect_stdout
        from io import StringIO
        # Lifespan banners must not corrupt the single safe JSON record.
        with redirect_stdout(StringIO()):
            result = verify(json.loads(os.environ['PEPPERYN_ISOLATION_BOOTSTRAP']),
                            os.environ['PEPPERYN_ISOLATION_ANON_KEY'],
                            os.environ['PEPPERYN_ISOLATION_SERVICE_KEY'], create_client,
                            live_transport=os.environ.get('PEPPERYN_OUTPUT_TRANSPORT') == 'A30_UVICORN')
    except Exception:
        result = dict(status='REFUSED', stage='LOCAL_INPUT_OR_IMPORT')
    print(json.dumps(result))
    raise SystemExit(0 if result['status'] in ('BOUNDED_A29_ASGI_OUTPUT_ISOLATION_PASS', 'BOUNDED_A30_UVICORN_OUTPUT_ISOLATION_PASS') else 1)
